"""Tests for src/report_generator.py"""

import os
import pandas as pd
from src.report_generator import ReportGenerator


def _build_sample_report(output_dir) -> ReportGenerator:
    report = ReportGenerator(output_dir=str(output_dir))
    report.add_title("Test Report")
    report.add_section("Intro", "This is a test section.")
    report.add_dict_as_table("Info", {"Rows": 100, "Columns": 5})
    report.add_dataframe("Sample Table", pd.DataFrame({"a": [1, 2], "b": [3, 4]}))
    report.add_insights("Insights", ["First insight.", "Second insight."])
    return report


def test_save_markdown_creates_file(tmp_path):
    report = _build_sample_report(tmp_path)
    path = report.save_markdown("test.md")

    assert os.path.exists(path)
    content = open(path).read()
    assert "Test Report" in content
    assert "First insight." in content


def test_save_pdf_creates_valid_pdf(tmp_path):
    report = _build_sample_report(tmp_path)
    path = report.save_pdf("test.pdf")

    assert os.path.exists(path)
    with open(path, "rb") as f:
        header = f.read(5)
    assert header == b"%PDF-"


def test_save_pdf_handles_nan_values(tmp_path):
    report = ReportGenerator(output_dir=str(tmp_path))
    report.add_title("NaN Test")
    report.add_dataframe("Table With NaN", pd.DataFrame({"a": [1, None], "b": [None, 2]}))

    path = report.save_pdf("nan_test.pdf")
    assert os.path.exists(path)


def test_save_excel_creates_sheets(tmp_path):
    report = _build_sample_report(tmp_path)
    path = report.save_excel("test.xlsx")

    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert "Summary" in wb.sheetnames
    assert any("Info" in name for name in wb.sheetnames)


def test_save_csv_summary_combines_tables(tmp_path):
    report = _build_sample_report(tmp_path)
    path = report.save_csv_summary("test_summary.csv")

    combined = pd.read_csv(path)
    assert "_section" in combined.columns
    assert set(combined["_section"].unique()) == {"Info", "Sample Table"}


def test_dict_as_table_preserves_integer_formatting(tmp_path):
    report = ReportGenerator(output_dir=str(tmp_path))
    report.add_dict_as_table("Numbers", {"Rows": 5001, "Ratio": 0.649})

    table_section = next(s for s in report.sections if s["heading"] == "Numbers")
    values = table_section["content"]["Value"].tolist()
    assert "5,001" in values
